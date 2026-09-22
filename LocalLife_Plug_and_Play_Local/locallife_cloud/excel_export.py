"""The simple operator workbook, LocalLife_Measurements.xlsx.

Built on every download from the canonical comparison CSV (paired_events.py),
which stays the detailed raw record. Sheets: RealSense, Logitech (one row per
finalised camera measurement) and Camera Comparison (one row per deposited
object). Nothing is computed that the raw CSV does not already hold.
"""

from __future__ import annotations

import io
from datetime import datetime
from typing import Any

CAMERA_COLUMNS = [
    "Object Number", "Object ID", "Date and Time", "Object Type", "Colour", "Material",
    "Sorting Result", "Geometry", "Dimensions (mm)", "Volume (L)", "Confidence", "Status",
]
COMPARISON_COLUMNS = [
    "Comparison Event ID", "Object Number", "Date and Time", "Object Type",
    "RealSense Volume (L)", "Logitech Volume (L)", "Absolute Difference (L)", "Percentage Difference",
    "RealSense Dimensions", "Logitech Dimensions", "RealSense Processing Time (ms)",
    "Logitech Processing Time (ms)", "Pairing Status",
]


def _number(value: Any) -> float | None:
    try:
        return None if value in (None, "") else float(value)
    except (TypeError, ValueError):
        return None


def _when(row: dict[str, Any]) -> datetime | None:
    """Local date and time from the row's ISO-8601 UTC value, or its epoch seconds."""
    text = row.get("timestamp_iso")
    if text:
        try:
            return datetime.fromisoformat(str(text)).astimezone().replace(tzinfo=None)
        except ValueError:
            pass
    stamp = _number(row.get("timestamp"))
    return None if stamp is None else datetime.fromtimestamp(stamp)


def dimensions_text(row: dict[str, Any]) -> str:
    """Cylinders as D x D x H (fitted), everything else as L x W x H."""
    diameter = _number(row.get("cylinder_diameter_mm"))
    height = _number(row.get("height_mm"))
    if row.get("geometry_method") == "cylinder" and diameter is not None and height is not None:
        return f"D{diameter:.0f} × D{diameter:.0f} × H{height:.0f}"
    values = [_number(row.get(key)) for key in ("length_mm", "width_mm", "height_mm")]
    return "" if None in values else " × ".join(f"{value:.0f}" for value in values)


def _camera_row(number: int, row: dict[str, Any]) -> list[Any]:
    confidence = _number(row.get("overall_confidence")) or _number(row.get("geometry_confidence"))
    status = row.get("status") or ""
    if row.get("reason"):
        status += f" ({row['reason']})"
    return [
        number, row.get("measurement_id"), _when(row), row.get("object_type"),
        row.get("colour"), row.get("material"), row.get("sorting_result"), row.get("geometry_method"),
        dimensions_text(row), _number(row.get("selected_volume_litres")),
        None if confidence is None else round(confidence, 3), status,
    ]


def build_workbook(rows: list[dict[str, Any]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheets = {"realsense": workbook.active, "logitech": workbook.create_sheet()}
    sheets["realsense"].title, sheets["logitech"].title = "RealSense", "Logitech"
    comparison = workbook.create_sheet("Camera Comparison")
    for sheet in (*sheets.values(), comparison):
        sheet.append(COMPARISON_COLUMNS if sheet is comparison else CAMERA_COLUMNS)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"

    counters = {"realsense": 0, "logitech": 0}
    events: dict[str, dict[str, dict[str, Any]]] = {}
    for row in rows:
        camera = row.get("camera_source")
        events.setdefault(row.get("comparison_event_id") or "", {})[camera] = row
        # A missing placeholder is a pairing outcome, not a measurement.
        if camera in sheets and row.get("status") != "missing":
            counters[camera] += 1
            sheets[camera].append(_camera_row(counters[camera], row))

    for number, (event_id, cameras) in enumerate(events.items(), start=1):
        realsense, logitech = cameras.get("realsense"), cameras.get("logitech")
        first = realsense or logitech or {}
        left = None if realsense is None else _number(realsense.get("selected_volume_litres"))
        right = None if logitech is None else _number(logitech.get("selected_volume_litres"))
        difference = None if left is None or right is None else abs(right - left)
        missing = [name for name, row in (("RealSense", realsense), ("Logitech", logitech))
                   if row is None or row.get("status") == "missing"]
        comparison.append([
            event_id, number, _when(first), first.get("object_type"),
            left, right,
            None if difference is None else round(difference, 6),
            None if difference is None or not left else round(difference / left * 100.0, 2),
            "" if realsense is None else dimensions_text(realsense),
            "" if logitech is None else dimensions_text(logitech),
            None if realsense is None else _number(realsense.get("processing_time_ms")),
            None if logitech is None else _number(logitech.get("processing_time_ms")),
            "paired" if not missing else "missing: " + ", ".join(missing),
        ])

    for sheet in (*sheets.values(), comparison):
        for index, column in enumerate(sheet.iter_cols(min_row=1, max_row=1), start=1):
            sheet.column_dimensions[get_column_letter(index)].width = max(12, len(str(column[0].value)) + 2)
        for cell in sheet["C"][1:]:
            cell.number_format = "yyyy-mm-dd hh:mm:ss"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
