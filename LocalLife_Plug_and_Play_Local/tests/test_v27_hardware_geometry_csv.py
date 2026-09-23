"""V27 regressions: fitted cylinders end to end, Logitech object masks, one CSV row per object.

These use deterministic synthetic depth and masks. They check the code path,
not the physical cameras: they are NOT hardware validation.
"""

from __future__ import annotations

import csv
import io
import json
from datetime import datetime
import math
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import numpy as np

from locallife_cloud.comparison import DualCameraCoordinator
from locallife_cloud.config import AppConfig
from locallife_cloud.logitech import bound_logitech_detections
from locallife_cloud.logitech_calibration import RELATIVE_ONLY_MESSAGE
from locallife_cloud.server import create_app
from locallife_cloud.shape_geometry import CUBOID, CYLINDER, IRREGULAR_RIGID, measure_shape
from locallife_cloud.types import CameraIntrinsics, Detection

RADIUS_M = 0.0415
HEIGHT_M = 0.195


def _bottle_shell(span_deg: float = 170.0, noise: float = 0.0015, seed: int = 11):
    """Front shell of an upright bottle seen obliquely, plus stray elevated pixels."""
    rng = np.random.default_rng(seed)
    angles = rng.uniform(math.radians(90 - span_deg / 2), math.radians(90 + span_deg / 2), 6000)
    points = np.column_stack((RADIUS_M * np.cos(angles), RADIUS_M * np.sin(angles))) + rng.normal(0, noise, (6000, 2))
    heights = rng.uniform(0.015, HEIGHT_M, 6000)
    strays = rng.uniform(-0.06, 0.06, (150, 2))
    return np.vstack([points, strays]), np.r_[heights, rng.uniform(0.015, 0.05, 150)]


class CylinderFitTests(unittest.TestCase):
    def test_upright_bottle_shell_reports_fitted_diameter_not_its_visible_arc(self) -> None:
        result = measure_shape(*_bottle_shell(), mesh_volume_l=0.6)
        self.assertEqual(result.geometry_method, CYLINDER)
        self.assertEqual(result.cylinder_orientation, "upright")
        self.assertAlmostEqual(result.cylinder_diameter_mm, 83, delta=3)
        self.assertEqual(result.length_mm, result.width_mm)  # diameter x diameter x height
        self.assertAlmostEqual(result.height_mm, 195, delta=3)
        self.assertAlmostEqual(result.radius_mm * 2, result.cylinder_diameter_mm)
        expected = math.pi * (result.radius_mm / 1000) ** 2 * (result.height_mm / 1000) * 1000
        self.assertAlmostEqual(result.selected_volume_litres, expected, places=6)
        self.assertGreater(result.fit_confidence, 0.5)
        data = result.to_dict()
        self.assertEqual(data["volume_liters"], data["selected_volume_litres"])
        self.assertEqual(data["diameter_mm"], data["cylinder_diameter_mm"])

    def test_horizontal_bottle_uses_axis_length_and_cross_section_circle(self) -> None:
        x, y = np.meshgrid(np.arange(-0.0975, 0.0975, 0.002), np.arange(-RADIUS_M, RADIUS_M, 0.002))
        points = np.column_stack((x.ravel(), y.ravel()))
        heights = RADIUS_M + np.sqrt(np.clip(RADIUS_M ** 2 - points[:, 1] ** 2, 0, None))
        heights += np.random.default_rng(3).normal(0, 0.0015, len(points))
        result = measure_shape(points, heights, mesh_volume_l=0.9)
        self.assertEqual(result.geometry_method, CYLINDER)
        self.assertEqual(result.cylinder_orientation, "lying")
        self.assertAlmostEqual(result.cylinder_diameter_mm, 83, delta=5)
        self.assertAlmostEqual(result.cylinder_height_mm, 195, delta=5)
        self.assertAlmostEqual(result.selected_volume_litres, math.pi * RADIUS_M ** 2 * 0.195 * 1000, delta=0.08)

    def test_poor_cylinder_fit_keeps_the_safe_height_map_result(self) -> None:
        # Only a 60 degree sliver of the shell is visible: too little arc to trust.
        points, heights = _bottle_shell(span_deg=60.0, noise=0.001)
        result = measure_shape(points, heights, mesh_volume_l=0.6, label="bottle")
        self.assertEqual(result.geometry_method, IRREGULAR_RIGID)
        self.assertEqual(result.selected_volume_litres, 0.6)
        self.assertIsNone(result.cylinder_diameter_mm)
        self.assertEqual(result.rejection_reason, "insufficient_arc_coverage")


class SharedDetector:
    device = "cpu"
    runtime = {"device": "cpu"}

    def __init__(self) -> None:
        self.items: list[Detection] = []

    def detect_batch(self, frames):
        return [[Detection(item.label, item.confidence, item.box, item.mask.copy(), color=item.color)
                 for item in self.items] for _ in frames]


class NoDepth:
    def estimate_batch(self, frames):
        return [np.full(frame.shape[:2], 1.0, dtype=np.float32) for frame in frames]


def _can_scene(size: int = 160, fx: float = 400.0, floor_m: float = 1.0, radius_m: float = 0.04,
               height_m: float = 0.15):
    """Overhead RealSense view of an upright can: a flat disc top above the floor."""
    camera = CameraIntrinsics(fx=fx, fy=fx, ppx=size / 2, ppy=size / 2, width=size, height=size)
    floor = np.full((size, size), floor_m, dtype=np.float32)
    rows, columns = np.mgrid[0:size, 0:size]
    top = floor_m - height_m
    x = (columns - size / 2) * top / fx
    y = (rows - size / 2) * top / fx
    disc = np.hypot(x, y) <= radius_m
    depth = floor.copy()
    depth[disc] = top
    frame = np.zeros((size, size, 3), dtype=np.uint8)
    frame[disc] = (30, 30, 200)
    ys, xs = np.nonzero(disc)
    detection = Detection("soda can", 0.9, (int(xs.min()), int(ys.min()), int(xs.max()) + 1, int(ys.max()) + 1),
                          disc, color="red")
    return camera, floor, depth, frame, detection


class RealSenseCylinderEndToEndTests(unittest.TestCase):
    def _station(self, directory: str, **overrides):
        detector = SharedDetector()
        settings = dict(
            results_dir=Path(directory), roi=(0, 0, 1, 1), min_component_pixels=20,
            tracker_confirm_frames=1, settle_frames=3, volume_window_frames=2,
            operating_mode="geometry_validation", auto_deposit=False, research_mode="realsense_only",
            enable_monocular_depth=False,
        )
        settings.update(overrides)
        manager = DualCameraCoordinator(AppConfig(**settings), detector=detector, depth_estimator=NoDepth())
        return manager, detector

    def _run(self, manager, detector, frames: int = 8):
        camera, floor, depth, frame, can = _can_scene()
        station = manager.camera("realsense")
        station.process_frame(np.zeros_like(frame), depth_m=floor, intrinsics=camera, persist=False)
        station.set_baseline()
        detector.items = [can]
        result = None
        for index in range(frames):
            result = station.process_frame(frame, depth_m=depth, intrinsics=camera, timestamp=100.0 + index)
        return station, result

    def test_one_finalised_can_is_one_row_with_matching_overlay_api_and_csv(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            manager, detector = self._station(directory)
            station, result = self._run(manager, detector)
            detection = result.detections[0]
            shape = detection.shape_geometry
            self.assertEqual(shape.geometry_method, CYLINDER)
            self.assertAlmostEqual(shape.cylinder_diameter_mm, 80, delta=5)
            self.assertAlmostEqual(shape.height_mm, 150, delta=5)
            # Every downstream field carries the fitted cylinder, not the footprint box.
            self.assertEqual(detection.footprint_length_mm, detection.footprint_width_mm)
            self.assertAlmostEqual(detection.footprint_length_mm, shape.cylinder_diameter_mm, places=1)
            self.assertAlmostEqual(detection.realsense_volume_l, shape.selected_volume_litres, places=5)
            api = create_app(manager.config, pipeline=manager).test_client()
            state_detection = api.get("/api/cameras/realsense/state").get_json()["latest"]["detections"][0]
            self.assertEqual(state_detection["shape_geometry"]["geometry_method"], CYLINDER)
            self.assertAlmostEqual(state_detection["realsense_volume_l"], shape.selected_volume_litres, places=5)

            response = api.get("/api/comparison/measurements.csv")
            self.assertEqual(response.headers["Cache-Control"], "no-store")
            rows = [row for row in csv.DictReader(io.StringIO(response.get_data(as_text=True).lstrip("﻿")))
                    if any(value not in ("", None) for value in row.values())]
            self.assertEqual(len(rows), 1)
            row = rows[0]
            self.assertEqual(row["geometry_method"], CYLINDER)
            self.assertAlmostEqual(float(row["diameter_mm"]), shape.cylinder_diameter_mm, places=1)
            self.assertAlmostEqual(float(row["volume_liters"]), shape.selected_volume_litres, places=5)
            for column in ("timestamp", "session_id", "comparison_event_id", "measurement_id", "camera",
                           "object_type", "colour", "height_mm", "status"):
                self.assertNotEqual(row[column], "", column)
            per_camera = list(csv.DictReader(io.StringIO(
                api.get("/api/cameras/realsense/measurements.csv").get_data(as_text=True).lstrip("﻿"))))
            self.assertEqual(len(per_camera), 1)
            self.assertEqual(per_camera[0]["event_id"], row["measurement_id"])

    def test_csv_failure_keeps_tracking_and_retry_writes_exactly_once(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            manager, detector = self._station(directory)
            station = manager.camera("realsense")
            with mock.patch.object(type(station.event_log), "_append", side_effect=OSError("disk full")):
                _, result = self._run(manager, detector)
            self.assertEqual(result.detections[0].track_id, 1)
            self.assertEqual(station.event_log.status()["persistence_failures"], 1)
            self.assertEqual(station.event_log.retry_failed()["recovered"], 1)
            self.assertEqual(station.event_log.retry_failed()["recovered"], 0)
            self.assertEqual(len(station.event_log.rows()), 1)

    def test_unsettled_object_is_recorded_once_as_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            manager, detector = self._station(directory, finalise_max_frames=4, settle_volume_tolerance=0.0)
            station = manager.camera("realsense")
            station._is_settled = lambda track_id: False  # a volume that never stabilises
            self._run(manager, detector, frames=10)
            rows = station.event_log.rows()
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]["status"], "rejected")
            self.assertEqual(rows[0]["reason"], "unstable_volume")

    def test_hardware_diagnostic_bundle_is_written_for_a_measurement(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            manager, detector = self._station(directory, hardware_diagnostic=True)
            self._run(manager, detector)
            bundles = list((Path(directory) / "realsense" / "hardware_diagnostics" / "realsense").glob("*_measurement_*"))
            self.assertEqual(len(bundles), 1)
            names = {path.name for path in bundles[0].iterdir()}
            self.assertTrue({"original.png", "object_mask.png", "depth.png", "points.npy", "result.json"} <= names)
            result = json.loads((bundles[0] / "result.json").read_text(encoding="utf-8"))
            self.assertEqual(result["geometry"]["geometry_method"], CYLINDER)
            self.assertIn("csv_row", result)
            scenes = list((Path(directory) / "realsense" / "hardware_diagnostics" / "realsense").glob("*_scene_*"))
            self.assertGreaterEqual(len(scenes), 1)


class LogitechObjectMaskTests(unittest.TestCase):
    def setUp(self) -> None:
        rng = np.random.default_rng(5)
        self.baseline = rng.integers(60, 190, (120, 160, 3), dtype=np.uint8)  # textured floor and sofa
        self.region = np.ones((120, 160), dtype=bool)
        self.object = np.zeros((120, 160), dtype=bool)
        self.object[40:80, 60:100] = True
        self.frame = self.baseline.copy()
        self.frame[self.object] = (40, 85, 145)
        self.shadow = np.zeros_like(self.object)
        self.shadow[80:95, 60:100] = True
        self.frame[self.shadow] = (self.baseline[self.shadow] * 0.6).astype(np.uint8)
        self.whole_roi = Detection("backpack", 0.8, (0, 0, 160, 120), np.ones((120, 160), dtype=bool))

    def test_whole_roi_detector_mask_is_cut_to_the_deposited_object(self) -> None:
        debug: dict = {}
        kept, _ = bound_logitech_detections(self.frame, self.baseline, [self.whole_roi], self.region,
                                            min_pixels=20, debug=debug)
        self.assertEqual(len(kept), 1)
        mask = kept[0].mask
        self.assertGreater((mask & self.object).sum() / self.object.sum(), 0.9)
        self.assertLess((mask & ~self.object).sum(), 0.1 * self.object.sum())
        self.assertLess((mask & self.shadow).sum(), 0.2 * self.shadow.sum())
        self.assertTrue(debug["final_mask_valid"])
        for key in ("detector", "foreground", "final", "rejected"):
            self.assertEqual(debug[key].shape, self.object.shape)
        self.assertGreater(debug["rejected"].sum(), 0.8 * (~self.object).sum())

    def test_no_empty_reference_means_no_object_mask(self) -> None:
        debug: dict = {}
        kept, warnings = bound_logitech_detections(self.frame, None, [self.whole_roi], self.region,
                                                   min_pixels=20, debug=debug)
        self.assertEqual(kept, [])
        self.assertIn("missing_empty_reference", debug["reasons"])
        self.assertTrue(warnings)

    def test_uncalibrated_logitech_says_calibration_required(self) -> None:
        self.assertEqual(RELATIVE_ONLY_MESSAGE, "CALIBRATION REQUIRED — metric volume unavailable")
        with tempfile.TemporaryDirectory() as directory:
            config = AppConfig(results_dir=Path(directory), roi=(0, 0, 1, 1), logitech_reference_distance_m=0.0)
            manager = DualCameraCoordinator(config, detector=SharedDetector(), depth_estimator=NoDepth())
            camera = CameraIntrinsics(fx=100, fy=100, ppx=80, ppy=60, width=160, height=120)
            logitech = manager.camera("logitech")
            logitech.process_frame(self.baseline, intrinsics=camera, persist=False)
            logitech.set_baseline()
            chain = logitech.logitech_diagnostics()
            self.assertTrue(chain["da_v2_loaded"])
            self.assertFalse(chain["calibration_loaded"])
            self.assertIn("CALIBRATION REQUIRED", logitech.logitech_calibration_status()["message"])


if __name__ == "__main__":
    unittest.main()


class MetricDepthStub:
    """Stand-in for Depth Anything V2 metric output: brighter pixels are closer."""

    device = "cpu"

    def estimate_batch(self, frames):
        return [2.0 - frame.max(axis=2).astype(np.float32) * 0.0015 for frame in frames]


class LogitechEndToEndTests(unittest.TestCase):
    """Calibrate the empty Logitech scene, then one bag reaches history, pairing and the workbook."""

    def _run(self, directory: str, logitech_frames: int = 4):
        detector = SharedDetector()
        config = AppConfig(
            results_dir=Path(directory), roi=(0, 0, 1, 1), min_component_pixels=20,
            tracker_confirm_frames=1, settle_frames=2, volume_window_frames=2,
            auto_deposit=True, logitech_reference_distance_m=0.0,
        )
        manager = DualCameraCoordinator(config, detector=detector, depth_estimator=MetricDepthStub())
        empty = np.zeros((40, 40, 3), dtype=np.uint8)
        floor = np.full((40, 40), 2.0, dtype=np.float32)
        camera = CameraIntrinsics(fx=100, fy=100, ppx=20, ppy=20, width=40, height=40)
        manager.camera("realsense").process_frame(empty, depth_m=floor, intrinsics=camera, persist=False)
        manager.camera("realsense").set_baseline()
        logitech = manager.camera("logitech")
        logitech.process_frame(empty, intrinsics=camera, persist=False)
        self.assertIn("CALIBRATION REQUIRED", logitech.logitech_calibration_status()["message"])
        logitech.calibrate_empty_scene(2.0)
        mask = np.zeros((40, 40), dtype=bool)
        mask[10:25, 10:25] = True
        frame = empty.copy()
        frame[mask] = (200, 0, 0)
        depth = floor.copy()
        depth[mask] = 1.7
        detector.items = [Detection("blue garbage bag", 0.9, (10, 10, 25, 25), mask, color="blue")]
        for index in range(4):
            manager.camera("realsense").process_frame(frame, depth_m=depth, intrinsics=camera, timestamp=100.0 + index)
        for index in range(logitech_frames):
            logitech.process_frame(frame, intrinsics=camera, timestamp=100.5 + index)
        return manager

    def test_calibrated_logitech_volume_reaches_history_pairing_and_workbook(self) -> None:
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as directory:
            manager = self._run(directory)
            logitech = manager.camera("logitech")
            status = logitech.logitech_calibration_status()
            self.assertTrue(status["metric_ready"])
            self.assertEqual(status["active_calibration"]["method"], "empty-plane-ray-alignment")
            self.assertTrue((Path(directory) / "logitech" / "calibration" / "logitech_depth.json").is_file())
            history = logitech.event_log.rows()
            self.assertEqual(len(history), 1)
            self.assertGreater(float(history[0]["volume_l"]), 0)
            rows = manager.paired_log.rows()
            self.assertEqual(sorted(row["camera_source"] for row in rows), ["logitech", "realsense"])
            self.assertEqual(len({row["comparison_event_id"] for row in rows}), 1)
            response = create_app(manager.config, pipeline=manager).test_client().get(
                "/api/export/LocalLife_Measurements.xlsx")
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.headers["Cache-Control"], "no-store")
            workbook = load_workbook(io.BytesIO(response.data))
            self.assertEqual(workbook.sheetnames, ["RealSense", "Logitech", "Camera Comparison"])
            logitech_rows = list(workbook["Logitech"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(len(logitech_rows), 1)
            self.assertGreater(logitech_rows[0][9], 0)  # Volume (L)
            self.assertEqual(len(list(workbook["RealSense"].iter_rows(min_row=2))), 1)
            pair = list(workbook["Camera Comparison"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(len(pair), 1)
            self.assertEqual(pair[0][-1], "paired")

    def test_late_logitech_result_fills_the_missing_pair(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            from locallife_cloud.paired_events import PairedComparisonLog

            log = PairedComparisonLog(Path(directory), window_seconds=10)
            event = log.record_measurement({"event_id": "rs-1", "camera_source": "realsense", "status": "accepted"},
                                           now=100)["comparison_event_id"]
            log.flush_expired(now=115)
            self.assertEqual([row["status"] for row in log.rows()], ["accepted", "missing"])
            late = log.record_measurement({"event_id": "lg-1", "camera_source": "logitech", "status": "accepted",
                                           "label": "clothing"}, now=118)
            self.assertEqual(late["comparison_event_id"], event)
            self.assertEqual([row["measurement_id"] for row in log.rows()], ["rs-1", "lg-1"])

    def test_uncalibrated_or_non_empty_scene_calibration_is_refused(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            manager = self._run(directory, logitech_frames=1)
            with self.assertRaises(ValueError):
                manager.camera("logitech").calibrate_empty_scene(2.0)  # the bag is still in view


class RealSenseCylinderPathTests(unittest.TestCase):
    def test_labelled_oblique_containers_select_cylinder_and_cuboid_stays_cuboid(self) -> None:
        rng = np.random.default_rng(1)

        def oblique(radius, height, n=5000):
            t, theta = rng.uniform(0, 1, n // 2), rng.uniform(0, 2 * math.pi, n // 2)
            top = np.column_stack((radius * np.sqrt(t) * np.cos(theta), radius * np.sqrt(t) * np.sin(theta)))
            arc = rng.uniform(math.pi * 0.1, math.pi * 0.9, n // 2)
            side = np.column_stack((radius * np.cos(arc), radius * np.sin(arc)))
            points = np.vstack([top, side]) + rng.normal(0, 0.001, (n, 2))
            heights = np.r_[np.full(n // 2, height), rng.uniform(0.01, height, n // 2)] + rng.normal(0, 0.001, n)
            return points, heights

        for label, radius, height in (("soda can", 0.033, 0.12), ("cream jar", 0.03, 0.05), ("bottle", 0.04, 0.2)):
            result = measure_shape(*oblique(radius, height), mesh_volume_l=math.pi * radius ** 2 * height * 1100,
                                   label=label)
            self.assertEqual(result.geometry_method, CYLINDER, label)
            self.assertAlmostEqual(result.cylinder_diameter_mm, radius * 2000, delta=radius * 2000 * 0.08)
            self.assertEqual(result.length_mm, result.width_mm)
        x, y = np.meshgrid(np.arange(-0.1, 0.1, 0.002), np.arange(-0.06, 0.06, 0.002))
        box = np.column_stack((x.ravel(), y.ravel()))
        self.assertEqual(measure_shape(box, 0.1 + rng.normal(0, 0.002, len(box)), label="bottle").geometry_method,
                         CUBOID)


class LogitechTrackingWithoutCalibrationTests(unittest.TestCase):
    """Detection and tracking must not wait for an empty reference or metric calibration."""

    def _station(self, directory: str):
        detector = SharedDetector()
        config = AppConfig(results_dir=Path(directory), roi=(0, 0, 1, 1), min_component_pixels=20,
                           tracker_confirm_frames=1, settle_frames=2, volume_window_frames=2,
                           operating_mode="geometry_validation", auto_deposit=False,
                           logitech_reference_distance_m=0.0, automatic_baseline=False)
        manager = DualCameraCoordinator(config, detector=detector, depth_estimator=MetricDepthStub())
        mask = np.zeros((60, 80), dtype=bool)
        mask[20:40, 30:50] = True
        frame = np.zeros((60, 80, 3), dtype=np.uint8)
        frame[mask] = (30, 30, 200)
        detector.items = [Detection("cream bottle", 0.6, (30, 20, 50, 40), mask, color="red")]
        return manager, frame

    def test_object_is_tracked_with_stable_id_before_any_baseline(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            manager, frame = self._station(directory)
            logitech = manager.camera("logitech")
            camera = CameraIntrinsics(fx=80, fy=80, ppx=40, ppy=30, width=80, height=60)
            ids = []
            for index in range(4):
                result = logitech.process_frame(frame, intrinsics=camera, timestamp=10.0 + index)
                ids.append(result.detections[0].track_id)
            self.assertIsNone(logitech.reference_rgb)
            self.assertEqual(len(set(ids)), 1)
            self.assertIsNotNone(ids[0])
            # V30 repair: without a baseline the object is still measured, as a
            # clearly labelled uncalibrated estimate rather than nothing at all.
            self.assertIsNotNone(result.detections[0].monocular_volume_l)
            self.assertEqual(result.detections[0].calibration_mode, "uncalibrated-estimate")
            report = logitech.stage_report()
            counters = report["counters"]
            self.assertEqual(counters["frames_processed"], 4)
            self.assertEqual(counters["raw_detections"], 4)
            self.assertEqual(counters["valid_masks"], 4)
            self.assertEqual(counters["detector_only_masks"], 4)
            self.assertGreaterEqual(counters["confirmed_tracks_total"], 1)
            self.assertEqual(counters["active_tracks_last_frame"], 1)
            state = logitech.state()
            self.assertIn("UNCALIBRATED ESTIMATE", state["volume_status"]["message"])
            # The RealSense station was never touched by any of this.
            self.assertEqual(manager.camera("realsense").stage_counters["frames_processed"], 0)

    def test_detector_diagnostic_saves_raw_and_final_predictions(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            manager, frame = self._station(directory)
            logitech = manager.camera("logitech")
            logitech.process_frame(frame, timestamp=1.0)
            response = create_app(manager.config, pipeline=manager).test_client().post(
                "/api/cameras/logitech/diagnose-detector")
            self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
            body = response.get_json()
            self.assertEqual(len(body["raw_predictions"]), 1)
            self.assertTrue(body["raw_predictions"][0]["passed_filters"])
            self.assertEqual(body["final_masks"], 1)
            saved = Path(body["saved_to"])
            for name in ("original.png", "raw_predictions.png", "final_masks.png", "relative_depth.png",
                         "summary.json"):
                self.assertTrue((saved / name).is_file(), name)

    def test_filter_rejection_reasons_are_counted(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            manager, frame = self._station(directory)
            logitech = manager.camera("logitech")
            logitech.config.logitech_detector_confidence = 0.9
            logitech.process_frame(frame, timestamp=1.0)
            report = logitech.stage_report()
            self.assertEqual(report["last_frame_rejections"], {"below_confidence": 1})
            self.assertIn("application filters", report["blocking_reason"])


class V28MetricCsvTests(unittest.TestCase):
    """V28: plane-aligned Logitech metric volume, valid rows, and a populated workbook."""

    def test_plane_alignment_recovers_scale_and_offset_from_the_empty_scene(self) -> None:
        from locallife_cloud.logitech_volume import fit_plane_alignment, ray_plane_distance

        camera = CameraIntrinsics(fx=500, fy=500, ppx=320, ppy=240, width=640, height=480)
        truth = ray_plane_distance((480, 640), camera, 1.60)
        predicted = (truth - 0.35) / 1.8  # a monocular map that is neither metric nor centred
        region = np.zeros((480, 640), dtype=bool)
        region[80:400, 120:520] = True
        calibration, diagnostics = fit_plane_alignment(predicted.astype(np.float32), region, camera, 1.60)
        self.assertIsNotNone(calibration)
        self.assertAlmostEqual(calibration.scale, 1.8, places=3)
        self.assertAlmostEqual(calibration.offset_m, 0.35, places=3)
        self.assertLess(diagnostics["plane_rmse_m"], 0.001)
        restored = calibration.apply(predicted.astype(np.float32))
        self.assertLess(float(np.abs(restored[region] - truth[region]).max()), 0.005)
        self.assertEqual(calibration.method, "empty-plane-ray-alignment")

    def test_height_map_volume_matches_a_known_block_and_filters_spikes(self) -> None:
        from locallife_cloud.logitech_volume import metric_object_volume, ray_plane_distance
        from locallife_cloud.volume import fit_reference_plane

        camera = CameraIntrinsics(fx=500, fy=500, ppx=160, ppy=120, width=320, height=240)
        floor = ray_plane_distance((240, 320), camera, 1.50).astype(np.float32)
        plane = fit_reference_plane(floor, camera)
        mask = np.zeros((240, 320), dtype=bool)
        mask[90:150, 110:190] = True  # 60 x 80 px block, 0.10 m tall
        depth = floor.copy()
        depth[mask] = (floor[mask] * (1.40 / 1.50))
        depth[100, 120] = 0.2  # a depth spike that would otherwise dominate
        result = metric_object_volume(depth, camera, mask, plane, min_height_m=0.01)
        self.assertIsNone(result.reason)
        # The occupied volume is the block's shadow on the support plane, so the
        # footprint is the mask's frustum cross-section AT THE PLANE (1.50 m),
        # not its smaller cross-section at the top face (1.40 m).
        footprint_m2 = 60 * 80 * (1.50 ** 2 / (500 * 500))
        expected = footprint_m2 * 0.10 * 1000
        self.assertAlmostEqual(result.measurement.liters, expected, delta=0.3 * expected)
        self.assertAlmostEqual(result.diagnostics["footprint_area_m2"], footprint_m2,
                               delta=0.15 * footprint_m2)
        self.assertGreaterEqual(result.diagnostics["rejected_spike_pixels"], 0)
        self.assertEqual(result.diagnostics["height_source"], "fitted_support_plane")
        for key in ("mask_pixels", "height_median_m", "footprint_area_m2", "footprint_cells",
                    "occlusion_filled_cells", "length_mm", "width_mm", "raw_volume_l"):
            self.assertIn(key, result.diagnostics)

    def test_stable_volume_waits_for_agreement(self) -> None:
        from locallife_cloud.logitech_volume import stable_volume

        self.assertIsNone(stable_volume([5.0, 9.0])[0])
        self.assertIsNone(stable_volume([5.0, 9.0, 14.0])[0])
        value, spread = stable_volume([8.0, 8.2, 8.1, 8.3, 8.15])
        self.assertAlmostEqual(value, 8.15, places=2)
        self.assertLess(spread, 0.2)

    def test_rows_always_carry_valid_time_and_finite_volume(self) -> None:
        from locallife_cloud.event_log import MeasurementEventLog

        with tempfile.TemporaryDirectory() as directory:
            log = MeasurementEventLog(Path(directory))
            log.record({"event_id": "a", "timestamp": None, "volume_l": float("nan"), "status": "accepted"})
            log.record({"event_id": "b", "timestamp": 1_700_000_000, "volume_l": 2.5, "status": "accepted"})
            rows = {row["event_id"]: row for row in log.rows()}
            self.assertTrue(rows["a"]["timestamp_iso"].endswith("+00:00"))
            self.assertEqual(rows["a"]["status"], "rejected")
            self.assertEqual(rows["a"]["reason"], "non_finite_volume")
            self.assertEqual(rows["a"]["volume_l"], "")
            self.assertEqual(rows["b"]["timestamp_iso"], "2023-11-14T22:13:20+00:00")
            self.assertEqual(float(rows["b"]["volume_l"]), 2.5)

    def test_a_soft_material_never_forces_box_geometry(self) -> None:
        from locallife_cloud.pipeline import classification_conflict

        cloth = Detection("cardboard box", 0.8, (0, 0, 10, 10), np.ones((10, 10), dtype=bool))
        cloth.material, cloth.material_confidence = "fabric or textile", 0.7
        self.assertIsNotNone(classification_conflict(cloth))
        box = Detection("cardboard box", 0.8, (0, 0, 10, 10), np.ones((10, 10), dtype=bool))
        box.material, box.material_confidence = "cardboard", 0.7
        self.assertIsNone(classification_conflict(box))

    def test_workbook_has_rows_for_both_cameras_and_the_pair(self) -> None:
        from openpyxl import load_workbook

        from locallife_cloud.paired_events import PairedComparisonLog

        with tempfile.TemporaryDirectory() as directory:
            log = PairedComparisonLog(Path(directory), session_id="s1")
            shape = {"geometry_method": "irregular_rigid", "selected_volume_litres": 3.0,
                     "length_mm": 300.0, "width_mm": 200.0, "height_mm": 100.0}
            log.record_measurement({"event_id": "rs-1", "camera_source": "realsense", "status": "accepted",
                                    "timestamp": 1_700_000_000, "label": "folded cloth", "colour": "blue",
                                    "material": "fabric or textile", "sorting_result": "allowed",
                                    "processing_time_ms": 40.0, "shape_geometry": shape}, now=100)
            log.record_measurement({"event_id": "lg-1", "camera_source": "logitech", "status": "accepted",
                                    "timestamp": 1_700_000_002, "label": "cardboard box",
                                    "classification_note": "detector says cardboard box; material says fabric",
                                    "processing_time_ms": 95.0,
                                    "shape_geometry": {**shape, "selected_volume_litres": 4.2}}, now=102)
            from locallife_cloud.excel_export import build_workbook

            workbook = load_workbook(io.BytesIO(build_workbook(log.rows())))
            realsense = list(workbook["RealSense"].iter_rows(min_row=2, values_only=True))
            logitech = list(workbook["Logitech"].iter_rows(min_row=2, values_only=True))
            comparison = list(workbook["Camera Comparison"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(len(realsense), 1)
            self.assertEqual(len(logitech), 1)
            self.assertEqual(len(comparison), 1)
            self.assertEqual(realsense[0][9], 3.0)
            self.assertEqual(logitech[0][9], 4.2)
            self.assertIsInstance(realsense[0][2], datetime)
            self.assertEqual(comparison[0][4], 3.0)
            self.assertEqual(comparison[0][5], 4.2)
            self.assertAlmostEqual(comparison[0][6], 1.2, places=6)
            self.assertAlmostEqual(comparison[0][7], 40.0, places=2)
            self.assertEqual(comparison[0][-1], "paired")
